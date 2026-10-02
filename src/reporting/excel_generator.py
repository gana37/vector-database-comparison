"""
Comprehensive Excel Report Generator for Vector Database Comparative Analysis.
Produces an executive-grade single-worksheet workbook meeting all technical specifications.
"""
import os
import json
from typing import Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.analysis.comparison import get_exhaustive_technical_data, DATABASES

# Styling Constants
FONT_FAMILY = "Segoe UI"

# Color Palette (Executive Navy / Corporate Palette)
COLOR_TITLE_BG = "1B365D"       # Deep Executive Navy
COLOR_SUBTITLE_BG = "2C3E50"    # Dark Slate
COLOR_META_BG = "EAEFF5"        # Soft Gray-Blue Tint
COLOR_HEADER_BG = "1B365D"      # Navy Table Header
COLOR_SECTION_BG = "243B53"     # Rich Slate Blue Section Divider
COLOR_ROW_ALT = "F8FAFC"        # Very Subtle Slate 50 Striping
COLOR_ROW_WHITE = "FFFFFF"      # Crisp White
COLOR_BORDER = "CBD5E1"         # Slate 300 Clean Border

# Type Badge Colors
COLOR_MEASURED_BG = "ECFDF5"    # Mint Green 50
COLOR_MEASURED_FG = "065F46"    # Forest Green 800
COLOR_DOC_BG = "EFF6FF"         # Soft Ice Blue 50
COLOR_DOC_FG = "1E40AF"         # Royal Blue 800

# Measured Status Colors
COLOR_BENCHMARKED_BG = "F0FDF4" # Subtle Green Tint
COLOR_UNBENCHMARKED_BG = "F8FAFC" # Muted Gray Tint
COLOR_UNBENCHMARKED_FG = "64748B" # Muted Gray Slate


def create_thin_border(color: str = COLOR_BORDER) -> Border:
    thin = Side(border_style="thin", color=color)
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def generate_excel_report(
    benchmark_summary: Optional[Dict[str, Any]] = None,
    output_path: str = "outputs/excel/VDB_Comprehensive_Analysis.xlsx"
) -> str:
    """
    Generates a single-worksheet Excel comparison workbook across 8 Vector Databases.
    
    Args:
        benchmark_summary: Dict containing measured benchmark metrics for each VDB.
        output_path: Target .xlsx file path.
        
    Returns:
        Absolute path to the generated workbook.
    """
    if benchmark_summary is None:
        summary_path = "outputs/raw_results/benchmark_summary.json"
        if os.path.exists(summary_path):
            with open(summary_path, "r", encoding="utf-8") as f:
                benchmark_summary = json.load(f)
        else:
            benchmark_summary = {}

    # Retrieve all comparison rows
    data_rows = get_exhaustive_technical_data(benchmark_summary)

    # Initialize openpyxl workbook
    wb = openpyxl.Workbook()
    
    # Strictly enforce ONE worksheet
    ws = wb.active
    ws.title = "VDB_Master_Comparison"

    # Thin border helper
    border_standard = create_thin_border()

    # =========================================================================
    # ROW 1: MASTER TITLE BANNER
    # =========================================================================
    ws.merge_cells("A1:K1")
    title_cell = ws["A1"]
    title_cell.value = "VECTOR DATABASE COMPARATIVE ANALYSIS AND BENCHMARKING REPORT"
    title_cell.font = Font(name=FONT_FAMILY, size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color=COLOR_TITLE_BG, end_color=COLOR_TITLE_BG, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    # Fill background for all merged cells in row 1
    for col_idx in range(1, 12):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = PatternFill(start_color=COLOR_TITLE_BG, end_color=COLOR_TITLE_BG, fill_type="solid")
        cell.border = border_standard

    # =========================================================================
    # ROW 2: SUBTITLE
    # =========================================================================
    ws.merge_cells("A2:K2")
    sub_cell = ws["A2"]
    sub_cell.value = "Comprehensive Architectural Evaluation, Empirical IR Benchmarks (BEIR/SciFact), and Production Trade-off Matrix across 8 Vector Engines"
    sub_cell.font = Font(name=FONT_FAMILY, size=10, italic=True, color="FFFFFF")
    sub_cell.fill = PatternFill(start_color=COLOR_SUBTITLE_BG, end_color=COLOR_SUBTITLE_BG, fill_type="solid")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    for col_idx in range(1, 12):
        cell = ws.cell(row=2, column=col_idx)
        cell.fill = PatternFill(start_color=COLOR_SUBTITLE_BG, end_color=COLOR_SUBTITLE_BG, fill_type="solid")
        cell.border = border_standard

    # =========================================================================
    # ROW 3: METADATA & EXPERIMENTAL SPECIFICATION
    # =========================================================================
    ws.merge_cells("A3:K3")
    meta_cell = ws["A3"]
    meta_cell.value = "Author: AI/ML Engineering Team | Target: AgentAnalytics.AI | Dataset: BEIR/SciFact (1,000 Docs / 50 Queries) | Model: all-MiniLM-L6-v2 (384-d, L2 Norm) | Workload: Top-K=10, 150 Executions"
    meta_cell.font = Font(name=FONT_FAMILY, size=9, bold=True, color="1E293B")
    meta_cell.fill = PatternFill(start_color=COLOR_META_BG, end_color=COLOR_META_BG, fill_type="solid")
    meta_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 20

    for col_idx in range(1, 12):
        cell = ws.cell(row=3, column=col_idx)
        cell.fill = PatternFill(start_color=COLOR_META_BG, end_color=COLOR_META_BG, fill_type="solid")
        cell.border = border_standard

    # =========================================================================
    # ROW 4: SPACER ROW
    # =========================================================================
    ws.row_dimensions[4].height = 8

    # =========================================================================
    # ROW 5: TABLE COLUMN HEADERS
    # =========================================================================
    headers = [
        "Technical Parameter",
        "Dimension / Category",
        "Data Type",
        "Qdrant",
        "Chroma",
        "FAISS",
        "Milvus",
        "Weaviate",
        "pgvector",
        "Elasticsearch",
        "Pinecone"
    ]
    ws.row_dimensions[5].height = 28
    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col_idx, value=header_text)
        cell.font = Font(name=FONT_FAMILY, size=10.5, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_standard

    # =========================================================================
    # ROWS 6+: DATA ROWS & SECTION DIVIDERS
    # =========================================================================
    current_row = 6
    current_section = None
    data_row_counter = 0

    for item in data_rows:
        section = item["section"]
        parameter = item["parameter"]
        data_type = item["type"]

        # Insert Section Divider if entering a new section
        if section != current_section:
            current_section = section
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=11)
            sec_cell = ws.cell(row=current_row, column=1)
            sec_cell.value = f"SECTION: {current_section.upper()}"
            sec_cell.font = Font(name=FONT_FAMILY, size=10.5, bold=True, color="FFFFFF")
            sec_cell.fill = PatternFill(start_color=COLOR_SECTION_BG, end_color=COLOR_SECTION_BG, fill_type="solid")
            sec_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws.row_dimensions[current_row].height = 24

            for col_idx in range(1, 12):
                c = ws.cell(row=current_row, column=col_idx)
                c.fill = PatternFill(start_color=COLOR_SECTION_BG, end_color=COLOR_SECTION_BG, fill_type="solid")
                c.border = border_standard

            current_row += 1

        # Populate Data Row
        data_row_counter += 1
        is_even = (data_row_counter % 2 == 0)
        row_bg = COLOR_ROW_ALT if is_even else COLOR_ROW_WHITE

        # Column 1: Parameter
        cell_param = ws.cell(row=current_row, column=1, value=parameter)
        cell_param.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="0F172A")
        cell_param.fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")
        cell_param.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        cell_param.border = border_standard

        # Column 2: Category
        # Shorten section name for category column
        cat_short = current_section.split(". ", 1)[-1] if ". " in current_section else current_section
        cell_cat = ws.cell(row=current_row, column=2, value=cat_short)
        cell_cat.font = Font(name=FONT_FAMILY, size=9, color="475569")
        cell_cat.fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")
        cell_cat.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell_cat.border = border_standard

        # Column 3: Data Type (MEASURED vs DOCUMENTATION)
        cell_type = ws.cell(row=current_row, column=3, value=data_type)
        if data_type == "MEASURED":
            cell_type.font = Font(name=FONT_FAMILY, size=8.5, bold=True, color=COLOR_MEASURED_FG)
            cell_type.fill = PatternFill(start_color=COLOR_MEASURED_BG, end_color=COLOR_MEASURED_BG, fill_type="solid")
        else:
            cell_type.font = Font(name=FONT_FAMILY, size=8.5, bold=True, color=COLOR_DOC_FG)
            cell_type.fill = PatternFill(start_color=COLOR_DOC_BG, end_color=COLOR_DOC_BG, fill_type="solid")
        cell_type.alignment = Alignment(horizontal="center", vertical="center")
        cell_type.border = border_standard

        # Columns 4 to 11: 8 Vector Databases
        for col_idx, db_name in enumerate(DATABASES, start=4):
            val = item.get(db_name, "N/A")
            cell_db = ws.cell(row=current_row, column=col_idx, value=val)
            cell_db.border = border_standard

            # Determine formatting based on data type and content
            if data_type == "MEASURED":
                is_unbenchmarked = "Not Benchmarked" in str(val) or "NOT BENCHMARKED" in str(val)
                if is_unbenchmarked:
                    cell_db.font = Font(name=FONT_FAMILY, size=9, italic=True, color=COLOR_UNBENCHMARKED_FG)
                    cell_db.fill = PatternFill(start_color=COLOR_UNBENCHMARKED_BG, end_color=COLOR_UNBENCHMARKED_BG, fill_type="solid")
                    cell_db.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                else:
                    # Successfully benchmarked measured value
                    cell_db.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="0F172A")
                    cell_db.fill = PatternFill(start_color=COLOR_BENCHMARKED_BG, end_color=COLOR_BENCHMARKED_BG, fill_type="solid")
                    cell_db.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                # Documentation content
                cell_db.font = Font(name=FONT_FAMILY, size=9, color="1E293B")
                cell_db.fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")
                cell_db.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # Dynamic row height for long text descriptions
        max_len = max(len(str(item.get(db, ""))) for db in DATABASES)
        if max_len > 120:
            ws.row_dimensions[current_row].height = 42
        elif max_len > 60:
            ws.row_dimensions[current_row].height = 32
        else:
            ws.row_dimensions[current_row].height = 22

        current_row += 1

    # =========================================================================
    # OPTIMAL COLUMN WIDTH CONFIGURATION
    # =========================================================================
    column_widths = {
        "A": 36,  # Technical Parameter
        "B": 24,  # Dimension / Category
        "C": 16,  # Data Type
        "D": 34,  # Qdrant
        "E": 34,  # Chroma
        "F": 34,  # FAISS
        "G": 34,  # Milvus
        "H": 34,  # Weaviate
        "I": 34,  # pgvector
        "J": 34,  # Elasticsearch
        "K": 34   # Pinecone
    }

    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # =========================================================================
    # FREEZE PANES
    # Freezes columns A-C and rows 1-5 so headers and parameters remain anchored
    # =========================================================================
    ws.freeze_panes = "D6"

    # Ensure output directory exists
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

    # Save workbook
    wb.save(output_path)

    # Strictly verify that there is only ONE sheet in the generated workbook
    verify_wb = openpyxl.load_workbook(output_path, read_only=True)
    sheet_count = len(verify_wb.sheetnames)
    verify_wb.close()

    if sheet_count != 1:
        raise ValueError(f"CRITICAL ERROR: Generated workbook has {sheet_count} sheets, expected exactly 1!")

    print(f"[EXCEL] Report generated successfully: {output_path}")
    print(f"[EXCEL] Verified exactly {sheet_count} worksheet: {ws.title}")
    print(f"[EXCEL] Total Rows: {current_row - 1} | Total Columns: 11")

    return os.path.abspath(output_path)


if __name__ == "__main__":
    generate_excel_report()
